"""공식 X 글(xlsx: 날짜, 링크, 본문 앞부분)을 일정에 맞춰 x.json 을 만든다.
사용: python3 tools/x_match.py 포레_트위터.xlsx events.csv [분류결과.xlsx]
- 그날(기간 포함) 일정이 하나면 그 일정에 붙인다.
- 여러 개면 일정명 낱말이 글에 가장 많이 나온 일정에 붙이고, 못 고르면 비워 둔다.
"""
import csv, json, re, sys, datetime as dt
import openpyxl

def days(a, b):
    a = dt.date.fromisoformat(a); b = dt.date.fromisoformat(b or a.isoformat())
    while a <= b:
        yield a.isoformat(); a += dt.timedelta(days=1)

def words(title):
    t = re.sub(r'\[[^\]]*\]|\([^)]*\)', ' ', title)
    return {w for w in re.split(r'[\s/·,:!?\'"<>《》「」]+', t.lower()) if len(w) >= 2}

def score(title, text):
    s = re.sub(r'\s', '', text.lower())
    return sum(1 for w in words(title) if w in s)

def main(xlsx, events_csv, out_xlsx=None):
    evs = {}
    for e in csv.DictReader(open(events_csv, encoding='utf-8')):
        if e.get('ID') and e.get('날짜') and e.get('공개여부', '공개') != '비공개':
            evs[e['ID']] = (e['날짜'][:10], (e.get('종료일') or '')[:10], e['일정명'])
    try:
        for e in json.load(open('cafe.json', encoding='utf-8'))['events']:
            evs.setdefault(e['id'], (e['date'], e.get('endDate') or '', e['title']))
    except FileNotFoundError:
        pass
    byday = {}
    for i, (a, b, t) in evs.items():
        for d in days(a, b):
            byday.setdefault(d, []).append(i)
    out, rows, seen = {}, [], set()
    ws = openpyxl.load_workbook(xlsx, read_only=True).worksheets[0]
    for r in ws.iter_rows(values_only=True):
        if not r or not r[0] or not r[1] or r[1] in seen:
            continue
        seen.add(r[1])
        d = r[0].date().isoformat() if hasattr(r[0], 'date') else str(r[0])[:10]
        text = str(r[2] or '')
        cand = byday.get(d, [])
        pick = None
        if len(cand) == 1:
            pick = cand[0]
        elif cand:
            sc = sorted(((score(evs[i][2], text), i) for i in cand), reverse=True)
            if sc[0][0] > 0 and sc[0][0] > sc[1][0]:
                pick = sc[0][1]
        if pick:
            out.setdefault(pick, []).append(r[1])
        rows.append((d, r[1], text, pick or '', evs[pick][2] if pick else ('일정 여러 개' if cand else '일정 없음')))
    for v in out.values():
        v.sort()  # 오래된 글 먼저
    json.dump({'source': '포레스텔라 공식 X @official_fore', 'events': out},
              open('x.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    if out_xlsx:
        wb = openpyxl.Workbook(); w = wb.active; w.title = 'X 분류'
        w.append(['날짜', '링크', '글', '일정 ID', '붙은 일정'])
        for x in rows: w.append(list(x))
        wb.save(out_xlsx)
    n = sum(len(v) for v in out.values())
    print(f'글 {len(rows)}개 중 {n}개를 일정 {len(out)}개에 붙임')

if __name__ == '__main__':
    main(*sys.argv[1:])
